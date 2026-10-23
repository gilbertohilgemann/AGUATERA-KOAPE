"""Exportación de un informe a Excel (.xlsx) o CSV, con la misma estructura que se ve en pantalla."""
import csv
import io
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

FORMATOS_EXCEL = {'gs': '#,##0', 'numero': '#,##0', 'decimal': '#,##0.00', 'porcentaje': '0.0', 'fecha': 'dd/mm/yyyy'}


def _valor(v, tipo):
    if isinstance(v, Decimal):
        return int(v) if tipo in ('gs', 'numero') and v == v.to_integral_value() else float(v)
    if tipo == 'fecha' and isinstance(v, str) and len(v) == 10:
        from datetime import date
        try:
            return date.fromisoformat(v)
        except ValueError:
            return v
    return v


def _seguro(texto, inicios=('=', '+', '-', '@', '\t', '\r')):
    """Evita la «inyección de fórmulas»: un TEXTO que empieza con = + - @ se toma literal. Los números no se tocan."""
    if isinstance(texto, str) and texto[:1] in inicios:
        return "'" + texto
    return texto


def a_excel(informe, datos, institucion, filtros_texto: str) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    libro = Workbook()
    hoja = libro.active
    hoja.title = informe.titulo[:31]
    hoja.append([institucion.nombre])
    hoja.append([informe.titulo])
    hoja.append([filtros_texto])
    hoja.append([f'Definición: {informe.definicion}'])
    hoja.append([f'Generado: {timezone.localtime():%d/%m/%Y %H:%M}'])
    hoja.append([])
    hoja['A1'].font = Font(bold=True, size=13)
    hoja['A2'].font = Font(bold=True)
    columnas = datos['columnas']
    hoja.append([c['titulo'] for c in columnas])
    fila_titulos = hoja.max_row
    for celda in hoja[fila_titulos]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill('solid', fgColor='EEEEEE')
    for fila in datos['filas']:
        # En .xlsx solo un texto que empieza con «=» se tomaría como fórmula.
        hoja.append([_seguro(_valor(fila.get(c['clave']), c['tipo']), ('=',)) for c in columnas])
    if datos.get('totales'):
        hoja.append([('Total' if n == 0 else _valor(datos['totales'].get(c['clave']), c['tipo'])
                      if c['clave'] in datos['totales'] else None) for n, c in enumerate(columnas)])
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)
    for n, c in enumerate(columnas, start=1):
        letra = get_column_letter(n)
        if c['tipo'] in FORMATOS_EXCEL:
            for fila in range(fila_titulos + 1, hoja.max_row + 1):
                hoja[f'{letra}{fila}'].number_format = FORMATOS_EXCEL[c['tipo']]
        hoja.column_dimensions[letra].width = max(12, min(45, len(c['titulo']) + 4))
    hoja.freeze_panes = hoja[f'A{fila_titulos + 1}']
    buffer = io.BytesIO()
    libro.save(buffer)
    respuesta = HttpResponse(buffer.getvalue(),
                             content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    respuesta['Content-Disposition'] = f'attachment; filename="{informe.codigo}-{timezone.localdate():%Y%m%d}.xlsx"'
    return respuesta


def a_csv(informe, datos) -> HttpResponse:
    """CSV con punto y coma y BOM UTF-8 (Excel en español lo abre bien). Números sin separador de miles."""
    buffer = io.StringIO()
    buffer.write('﻿')
    escritor = csv.writer(buffer, delimiter=';')
    columnas = datos['columnas']
    escritor.writerow([c['titulo'] for c in columnas])
    def celda(valor):
        if isinstance(valor, Decimal):
            return str(valor).replace('.', ',')  # número (también negativo): sin apóstrofo
        if isinstance(valor, (int, float)):
            return valor
        return _seguro(valor if valor is not None else '')

    for fila in datos['filas']:
        escritor.writerow([celda(fila.get(c['clave'])) for c in columnas])
    if datos.get('totales'):
        escritor.writerow(['Total' if n == 0 else celda(datos['totales'].get(c['clave'])) if c['clave'] in datos['totales']
                           else '' for n, c in enumerate(columnas)])
    respuesta = HttpResponse(buffer.getvalue(), content_type='text/csv; charset=utf-8')
    respuesta['Content-Disposition'] = f'attachment; filename="{informe.codigo}-{timezone.localdate():%Y%m%d}.csv"'
    return respuesta

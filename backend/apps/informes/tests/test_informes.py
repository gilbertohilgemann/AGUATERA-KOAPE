"""Fase 8: informes, panel, exportación y credenciales de integración (solo lectura)."""
import io
from decimal import Decimal as D

from django.core.cache import cache
from rest_framework.test import APIClient

from apps.caja.tests.test_caja import BaseCaja
from apps.informes.consultas import INFORMES
from apps.integraciones.models import RegistroSolicitudIntegracion
from comun.pruebas.utilidades import admin_de, cliente_de, nueva_institucion, usuario_con_permisos


class InformesTests(BaseCaja):
    def informe(self, codigo, api=None, codigo_http=200, **filtros):
        r = (api or self.api).get(f'/api/v1/informes/{codigo}/', filtros)
        self.assertEqual(r.status_code, codigo_http, getattr(r, 'data', r))
        return r.data

    def test_todos_los_informes_responden_con_su_definicion(self):
        for informe in INFORMES:
            datos = self.informe(informe.codigo)
            self.assertTrue(datos['definicion'])
            self.assertTrue(datos['columnas'], informe.codigo)
        self.assertEqual(len(self.api.get('/api/v1/informes/').data), len(INFORMES))

    def test_facturado_cobrado_y_saldos_no_se_confunden(self):
        f = self.credito([self.oct])                                  # 44.800 a crédito
        self.cobrar(20000, [(f, 20000)])
        r = self.cobrar(4800, [(f, 4800)])
        self.api.post(f'/api/v1/recibos/{r["numero"]}/anular/', {'motivo': 'error'}, format='json')  # revertido
        self.assertEqual(D(self.informe('facturas')['totales']['total']), D('44800'))
        self.assertEqual(D(self.informe('recaudacion')['totales']['total']), D('20000'))
        por_forma = self.informe('recaudacion', agrupar='forma')
        self.assertEqual([(x['grupo'], D(x['total'])) for x in por_forma['filas']], [('Efectivo', D('20000'))])
        saldos = self.informe('saldos_clientes')
        fila = saldos['filas'][0]
        self.assertEqual((fila['cliente'], D(fila['saldo'])), ('C1', D('24800')))
        self.assertEqual(D(fila['a_facturar']), self.nov.total + self.otra.total)   # consumo sin facturar aparte
        tramos = {x['tramo']: D(x['saldo']) for x in self.informe('antiguedad_deuda')['filas']}
        self.assertEqual(tramos['Por vencer'], D('24800'))
        self.assertEqual(D(self.informe('liquidaciones', periodo_desde='2026-01', periodo_hasta='2026-12')['totales']['total']),
                         self.oct.total + self.nov.total + self.otra.total)

    def test_caja_y_gastos(self):
        g = self.gasto()
        self.api.post(f'/api/v1/gastos/{g["referencia"]}/confirmar/')
        self.gasto(importe=5000)                                       # sin confirmar
        datos = self.informe('gastos')
        self.assertEqual((D(datos['totales']['total']), D(datos['totales']['sin_confirmar'])), (D('15000'), D('5000')))
        ref = self.sesion()['referencia']
        self.cerrar(ref, {'EFECTIVO': 80000}, 'faltan 5.000')
        cajas = self.informe('cajas', estado='en_revision')
        self.assertEqual([(x['sesion'], x['diferencias']) for x in cajas['filas']], [(ref, 'Efectivo -5.000')])

    def test_exportar_excel_y_csv(self):
        self.credito([self.oct])
        r = self.api.get('/api/v1/informes/saldos_clientes/', {'formato': 'xlsx'})
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        hoja = load_workbook(io.BytesIO(r.content)).active
        valores = [fila for fila in hoja.iter_rows(values_only=True)]
        titulos = valores.index(next(v for v in valores if v and v[0] == 'Cliente'))
        self.assertEqual(valores[titulos + 1][0], 'C1')
        self.assertIsInstance(valores[titulos + 1][3], (int, float))        # número real, se puede sumar en Excel
        r = self.api.get('/api/v1/informes/saldos_clientes/', {'formato': 'csv'})
        texto = r.content.decode('utf-8')
        self.assertTrue(texto.startswith('﻿Cliente;Nombre;'))
        self.assertEqual(self.api.get('/api/v1/informes/facturas/', {'formato': 'pdf'}).status_code, 400)

    def test_permisos_filtros_y_aislamiento(self):
        usuario_con_permisos(self.a, 'padron', ['padron.consultar'])
        padron = cliente_de(self.a, 'padron')
        self.assertEqual([i['codigo'] for i in padron.get('/api/v1/informes/').data], ['conexiones'])
        self.informe('recaudacion', api=padron, codigo_http=403)
        self.informe('recaudacion', api=padron, codigo_http=403, formato='xlsx')
        self.assertEqual([s['titulo'] for s in padron.get('/api/v1/informes/panel/').data], ['Padrón'])
        self.informe('facturas', codigo_http=400, desde='2026-13-01')
        self.informe('lecturas', codigo_http=400, periodo_desde='2026-1')
        self.informe('no-existe', codigo_http=404)
        self.credito([self.oct])
        b = nueva_institucion('JUNTA-B')
        admin_de(b)
        otra = cliente_de(b)
        self.assertEqual(self.informe('facturas', api=otra)['filas'], [])
        self.assertEqual(self.informe('conexiones', api=otra)['filas'], [])
        self.assertEqual(self.informe('conexiones')['totales']['total'], '2')

    def test_panel(self):
        self.emitir([self.oct])
        secciones = {s['titulo']: {i['titulo']: i for i in s['indicadores']} for s in self.api.get('/api/v1/informes/panel/').data}
        cobranzas = secciones['Facturación y cobranzas del mes']
        self.assertEqual((D(cobranzas['Facturado']['valor']), D(cobranzas['Cobrado']['valor'])), (D('44800'), D('44800')))
        self.assertTrue(cobranzas['Facturado']['definicion'])
        self.assertEqual(secciones['Padrón']['Conexiones']['valor'], '2')


class IntegracionTests(BaseCaja):
    def setUp(self):
        super().setUp()
        cache.clear()

    def credencial(self, permisos=('padron.consultar',), codigo=201, **extra):
        r = self.api.post('/api/v1/integraciones/credenciales/', {'nombre': 'ERP contable', 'permisos': list(permisos),
                                                                   **extra}, format='json')
        self.assertEqual(r.status_code, codigo, r.data)
        return r.data

    def externo(self, clave):
        api = APIClient()
        api.credentials(HTTP_AUTHORIZATION=f'Clave {clave}')
        return api

    def test_credencial_de_solo_lectura_con_sus_permisos(self):
        datos = self.credencial()
        self.assertTrue(datos['clave'].startswith(f'kpa_{datos["prefijo"]}_'))
        listado = self.api.get('/api/v1/integraciones/credenciales/').data
        self.assertNotIn('clave', listado[0])                          # la clave se muestra una sola vez
        erp = self.externo(datos['clave'])
        self.assertEqual(erp.get('/api/v1/clientes/').status_code, 200)
        self.assertEqual(erp.get('/api/v1/informes/conexiones/').status_code, 200)
        self.assertEqual(erp.get('/api/v1/facturas/').status_code, 403)              # sin ese permiso
        self.assertEqual(erp.post('/api/v1/clientes/', {'Code': 'X', 'nombre': 'x'}, format='json').status_code, 403)
        self.assertEqual(erp.get('/api/v1/auth/yo/').status_code, 403)
        self.assertEqual(erp.get('/api/v1/integraciones/credenciales/').status_code, 403)
        self.assertEqual(RegistroSolicitudIntegracion.objects.count(), 6)            # todas registradas
        registro = self.api.get(f'/api/v1/integraciones/credenciales/{datos["prefijo"]}/registro/').data
        self.assertEqual(registro[-1]['ruta'], '/api/v1/clientes/')

    def test_revocar_vencer_y_limitar(self):
        datos = self.credencial(limite_por_minuto=2)
        erp = self.externo(datos['clave'])
        self.assertEqual([erp.get('/api/v1/clientes/').status_code for _ in range(3)], [200, 200, 429])
        url = f'/api/v1/integraciones/credenciales/{datos["prefijo"]}/revocar/'
        self.assertEqual(self.api.post(url, {'motivo': ''}, format='json').status_code, 400)
        self.assertEqual(self.api.post(url, {'motivo': 'Se cambió de sistema'}, format='json').data['activa'], False)
        cache.clear()
        self.assertEqual(erp.get('/api/v1/clientes/').status_code, 401)
        self.assertEqual(self.externo('kpa_inventada_xxx').get('/api/v1/clientes/').status_code, 401)
        self.credencial(expira_en='2020-01-01T00:00:00Z', codigo=400)

    def test_permisos_validos_y_aislamiento(self):
        self.credencial(permisos=['cobranzas.registrar'], codigo=400)              # no es de consulta
        self.credencial(permisos=[], codigo=400)
        usuario_con_permisos(self.a, 'cajero', ['cobranzas.registrar'])
        self.assertEqual(cliente_de(self.a, 'cajero').get('/api/v1/integraciones/credenciales/').status_code, 403)
        b = nueva_institucion('JUNTA-B')
        admin_de(b)
        from apps.padron import servicios as padron
        padron.crear_cliente(institucion=b, datos={'Code': 'SOLO-B', 'nombre': 'De B'})
        datos = self.credencial()
        clientes = self.externo(datos['clave']).get('/api/v1/clientes/').data['results']
        self.assertEqual([c['Code'] for c in clientes], ['C1'])                     # solo su institución
        self.assertEqual(cliente_de(b).get(f'/api/v1/integraciones/credenciales/{datos["prefijo"]}/').status_code, 404)


class RevisionFase8Tests(BaseCaja):
    def setUp(self):
        super().setUp()
        cache.clear()

    def test_no_se_dan_permisos_que_el_creador_no_tiene(self):
        usuario_con_permisos(self.a, 'integrador', ['integraciones.administrar', 'padron.consultar'])
        integrador = cliente_de(self.a, 'integrador')
        r = integrador.post('/api/v1/integraciones/credenciales/', {'nombre': 'X', 'permisos': ['caja.supervisar']},
                            format='json')
        self.assertEqual(r.data['codigo'], 'permiso_ajeno')
        r = integrador.post('/api/v1/integraciones/credenciales/', {'nombre': 'X', 'permisos': ['padron.consultar']},
                            format='json')
        self.assertEqual(r.status_code, 201)

    def test_filtros_invalidos_no_rompen(self):
        for parametros in ({'zona': 'a\x00b'}, {'cobrador': 'x\x00'}, {'estado': 'otro'}, {'zona': 'x' * 200}):
            for codigo in ('conexiones', 'cajas', 'recaudacion'):
                self.assertEqual(self.api.get(f'/api/v1/informes/{codigo}/', parametros).status_code, 400, parametros)

    def test_panel_por_permiso_de_cada_indicador(self):
        usuario_con_permisos(self.a, 'cobros', ['cobranzas.consultar'])
        titulos = [i['titulo'] for s in cliente_de(self.a, 'cobros').get('/api/v1/informes/panel/').data
                   for i in s['indicadores']]
        self.assertIn('Cobrado', titulos)
        self.assertNotIn('Facturado', titulos)

    def test_csv_con_negativos_y_totales(self):
        self.emitir([self.oct])
        ref = self.sesion()['referencia']
        self.cerrar(ref, {'EFECTIVO': 144000}, 'faltan')
        texto = self.api.get('/api/v1/informes/cajas/', {'formato': 'csv'}).content.decode('utf-8')
        self.assertIn('Efectivo -800', texto)
        self.assertNotIn("'-", texto)
        self.assertIn('\nTotal;', texto.replace('\r\n', '\n'))

    def test_clave_revocada_queda_registrada(self):
        r = self.api.post('/api/v1/integraciones/credenciales/', {'nombre': 'ERP', 'permisos': ['padron.consultar']},
                          format='json').data
        self.api.post(f'/api/v1/integraciones/credenciales/{r["prefijo"]}/revocar/', {'motivo': 'x'}, format='json')
        api = APIClient()
        api.credentials(HTTP_AUTHORIZATION=f'Clave {r["clave"]}')
        self.assertEqual(api.get('/api/v1/clientes/').status_code, 401)
        self.assertEqual(RegistroSolicitudIntegracion.objects.get().estado, 401)

    def test_gasto_con_varios_adjuntos_se_cuenta_una_vez(self):
        from apps.caja.models import AdjuntoGasto, Gasto
        g = self.gasto()
        self.api.post(f'/api/v1/gastos/{g["referencia"]}/confirmar/')
        gasto = Gasto.objects.get()
        for n in range(3):
            AdjuntoGasto.objects.create(institucion=self.a, gasto=gasto, tipo='comprobante', archivo=f'x/{n}.png',
                                        subido_por=gasto.registrado_por)
        totales = self.api.get('/api/v1/informes/gastos/').data['totales']
        self.assertEqual((totales['confirmados'], D(totales['total'])), ('1', D('15000')))
